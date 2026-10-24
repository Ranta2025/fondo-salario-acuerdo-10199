"""Best-effort import of areas and positions from the legacy "Dist.Monto x area" sheet.

Layout: column B = area number, C = name, D = vacant slots (the count multiplied by salary and months
to get the unexecuted amount; the legacy header says "approved" but it is only filled for vacant positions),
E = salary, G = months unexecuted.
Area header rows have a name like "3.- Departamento"; the rows below with the same number are positions.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from fondo_salario.importers.workers import (
    _clean_text,
    _normalize_header,
    pick_table,
    roster_entries,
)

AREA_HEADER = re.compile(r"^\s*(\d+)\s*[.\-]")


@dataclass
class PositionRow:
    name: str
    vacant: int = 0
    salary: Decimal = Decimal(0)
    months_unexecuted: int = 0


@dataclass
class AreaRow:
    number: int
    name: str
    positions: list[PositionRow] = field(default_factory=list)


def _number(value: object) -> Decimal:
    try:
        return Decimal(str(value)) if value not in (None, "") else Decimal(0)
    except InvalidOperation:
        return Decimal(0)


def _pick_sheet(workbook):
    for sheet in workbook.worksheets:
        key = _normalize_header(sheet.title)
        if "monto" in key and "area" in key:
            return sheet
    return workbook.worksheets[0]


def _clean_area_name(text: str) -> str:
    return re.sub(r"^\s*\d+\s*[.\-]+\s*", "", text).strip()


def read_roster_structure(path: str | Path) -> list[AreaRow] | None:
    """Areas and positions from a staffing roster (one row per position; rows without a worker are vacant)."""
    picked = pick_table(Path(path))
    if picked is None or "roster" not in picked[3]:
        return None
    _, table, header, columns = picked
    areas: dict[str, AreaRow] = {}
    for _line, area, title, name, _national_id, _group in roster_entries(table, header, columns):
        if not area:
            continue
        current = areas.setdefault(area, AreaRow(len(areas) + 1, area))
        if title is None:
            continue
        position = next(
            (p for p in current.positions if _normalize_header(p.name) == _normalize_header(title)), None
        )
        if position is None:
            position = PositionRow(title)
            current.positions.append(position)
        position.vacant += int(not name)
    return list(areas.values())


def read_table_structure(path: str | Path) -> list[AreaRow] | None:
    """Areas and positions from a plain table (Area | Cargo | Vacantes | Salario | Meses), on any sheet."""
    picked = pick_table(Path(path), "structure")
    if picked is None:
        return None
    _, table, header, columns = picked
    areas: dict[str, AreaRow] = {}
    current = ""
    for raw in table[header + 1 :]:

        def cell(key: str, raw=raw) -> object:
            col = columns.get(key)
            return raw[col] if col is not None and col < len(raw) else None

        current = _clean_text(cell("area")) or current  # merged cells leave the area blank below the first
        if not current:
            continue
        row = areas.setdefault(current, AreaRow(len(areas) + 1, current))
        title = _clean_text(cell("position"))
        if not title:
            continue
        position = next(
            (p for p in row.positions if _normalize_header(p.name) == _normalize_header(title)), None
        )
        if position is None:
            position = PositionRow(title)
            row.positions.append(position)
        position.vacant += int(_number(cell("vacant")))
        position.salary = _number(cell("salary")) or position.salary
        position.months_unexecuted = int(_number(cell("months"))) or position.months_unexecuted
    return list(areas.values())


def _has_legacy_sheet(workbook) -> bool:
    return any(
        "monto" in _normalize_header(s.title) and "area" in _normalize_header(s.title)
        for s in workbook.worksheets
    )


def read_structure(path: str | Path, sheet: str | None = None) -> list[AreaRow]:
    """Legacy "Dist.Monto x area" sheet, else a staffing roster, else a plain areas/positions table."""
    if sheet is None and Path(path).suffix.lower() == ".csv":
        return read_table_structure(path) or []
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if sheet is None and not _has_legacy_sheet(workbook):
            for reader in (read_roster_structure, read_table_structure):
                found = reader(path)
                if found is not None:
                    return found
        worksheet = workbook[sheet] if sheet else _pick_sheet(workbook)
        areas: list[AreaRow] = []
        by_number: dict[int, AreaRow] = {}
        for row in worksheet.iter_rows(min_col=2, max_col=7, values_only=True):
            number_cell, name_cell, vacant, salary, _, months = row
            if not isinstance(number_cell, (int, float)) or not name_cell:
                continue
            number, name = int(number_cell), str(name_cell).strip()
            header = AREA_HEADER.match(name)
            if header and number not in by_number:
                area = AreaRow(number, _clean_area_name(name))
                areas.append(area)
                by_number[number] = area
                continue
            area = by_number.get(number)
            if area is not None:
                area.positions.append(
                    PositionRow(
                        name,
                        int(_number(vacant)),
                        _number(salary),
                        int(_number(months)),
                    )
                )
        return areas
    finally:
        workbook.close()


def looks_like(path: str | Path) -> str | None:
    """ "workers" or "structure" when a table of that kind is found in the file, else None."""
    try:
        if pick_table(Path(path), "workers") is not None:
            return "workers"
        if pick_table(Path(path), "structure") is not None:
            return "structure"
    except Exception:  # unreadable file: the caller reports the real error
        return None
    return None
