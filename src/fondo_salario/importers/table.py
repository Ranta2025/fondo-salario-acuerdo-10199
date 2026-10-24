"""The one fixed table structure used to import data, to export it and for the blank template.

Two sheets, with the same titles as the tables of the program:

* «Trabajadores»: the Trabajadores table.
* «Áreas y cargos»: one row per position (a row with only the area creates an empty area).

A file must have these titles in its first row, otherwise the import is refused. Importing only adds
records that do not exist yet; nothing that is already stored is changed or deleted.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fondo_salario.domain.models import Area, Position, Worker
from fondo_salario.domain.scales import parse_roman, to_roman

WORKERS_SHEET = "Trabajadores"
STRUCTURE_SHEET = "Áreas y cargos"
INSTRUCTIONS_SHEET = "Instrucciones"

WORKERS_COLUMNS = [
    "Nombre y apellidos",
    "Carné de identidad",
    "Área",
    "Cargo",
    "Grupo",
    "Antigüedad",
    "Días",
    "Contribución",
    "Derecho a cobro",
]
STRUCTURE_COLUMNS = [
    "Área",
    "Cargo",
    "Aprobadas",
    "Vacantes",
    "Temporales",
    "Vacantes a distribuir",
    "Salario",
    "Meses",
]
YES, NO = "Sí", "No"
_YES_WORDS = {"si", "s", "x", "1", "true", "yes"}
_NO_WORDS = {"no", "n", "0", "false", ""}


class ImportFormatError(Exception):
    """The file does not follow the fixed structure; the message says what is wrong."""


def key(text: object) -> str:
    """Comparable form of a name or title: no accents, case, spaces or punctuation."""
    ascii_text = unicodedata.normalize("NFKD", str(text if text is not None else "")).encode(
        "ascii", "ignore"
    )
    return re.sub(r"[^a-z0-9]", "", ascii_text.decode().lower())


def clean_id(value: object) -> str:
    """ID number as text; short all-digit numbers lost their leading zeros in Excel (IDs have 11 digits)."""
    text = clean(value)
    return text.zfill(11) if text.isdigit() and len(text) < 11 else text


def clean(value: object) -> str:
    if isinstance(value, float) and value == int(value):
        value = int(value)
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


# -- reading -----------------------------------------------------------------------------------
@dataclass
class WorkerRow:
    line: int
    name: str
    national_id: str
    area: str
    position: str
    salary_group: int
    seniority_years: int
    days_worked: int
    contribution: int
    eligible: bool


@dataclass
class StructureRow:
    line: int
    area: str
    position: str  # blank: the row only creates the area
    approved: int = 0
    vacant: int = 0
    temp_vacant: int = 0
    distribute: int | None = None
    salary: Decimal = Decimal(0)
    months_unexecuted: int = 0


@dataclass
class ReadResult:
    rows: list = field(default_factory=list)
    errors: list[str] = field(default_factory=list)  # «Fila N: …» rows that cannot be used


def _load(path: Path, sheet_name: str) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = {key(ws.title): ws for ws in workbook.worksheets}
            sheet = sheets.get(key(sheet_name)) or (workbook.worksheets[0] if len(sheets) == 1 else None)
            if sheet is None:
                raise ImportFormatError(
                    f"El archivo no tiene la hoja «{sheet_name}». "
                    "Use «Guardar plantilla» para ver la estructura."
                )
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".csv":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:  # CSV saved by Excel on Windows
            text = raw.decode("cp1252", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text, newline=""), dialect)]
    raise ImportFormatError("Solo se pueden importar archivos de Excel (.xlsx) o CSV.")


def _columns(table: list[list[object]], expected: list[str], what: str) -> dict[str, int]:
    """Column of every expected title; the titles must be in the first row."""
    header = {key(v): i for i, v in enumerate(table[0])} if table else {}
    missing = [title for title in expected if key(title) not in header]
    if missing:
        raise ImportFormatError(
            f"El archivo de {what} no sigue la estructura. Faltan estas columnas en la primera fila: "
            f"{', '.join(missing)}.\n\nLa primera fila debe tener estos títulos: {', '.join(expected)}.\n"
            "Use «Guardar plantilla» para obtener un archivo con la estructura correcta."
        )
    return {title: header[key(title)] for title in expected}


def _number(value: object, what: str, low: int, high: int, required: bool = True) -> int | None:
    text = clean(value)
    if not text:
        if required:
            raise ValueError(f"falta {what}")
        return None
    if not re.fullmatch(r"\d+([.,]0{1,2})?", text):
        raise ValueError(f"{what} debe ser un número entero sin separadores («{text}»)")
    number = int(re.split(r"[.,]", text)[0])
    if not low <= number <= high:
        raise ValueError(f"{what} debe ser un entero entre {low} y {high} («{text}»)")
    return int(number)


def _group(value: object) -> int:
    text = clean(value)
    if not text:
        raise ValueError("falta el grupo")
    number = int(text) if text.isdigit() else parse_roman(text)
    if number is None or not 1 <= number <= 25:
        raise ValueError(f"el grupo debe ser un número romano de I a XXV («{text}»)")
    return number


def _yes_no(value: object) -> bool:
    text = key(value)
    if text in _YES_WORDS:
        return True
    if text in _NO_WORDS:
        return False
    raise ValueError(f"«Derecho a cobro» debe ser Sí o No («{clean(value)}»)")


def _money(value: object) -> Decimal:
    """Amount; a lone separator followed by exactly three digits is thousands, otherwise decimals."""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        text = str(value)  # a real number cell has no separator ambiguity: 1.234 is one and a bit
    else:
        text = clean(value).replace(" ", "")
        if not text:
            return Decimal(0)
        text = _text_amount(text)
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"el salario debe ser un número («{text}»)") from None
    if not amount.is_finite() or amount < 0:
        raise ValueError(f"el salario debe ser un número positivo («{text}»)")
    return amount


def _text_amount(text: str) -> str:
    """Typed amount as a plain decimal string."""
    if "," in text and "." in text:  # the last one is the decimal separator
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        return text.replace(thousands, "").replace(decimal, ".")
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", text):
        return re.sub(r"[.,]", "", text)
    return text.replace(",", ".")


def read_workers(path: str | Path) -> ReadResult:
    """Rows of the «Trabajadores» table; rows that break the rules go to `errors`."""
    table = _load(Path(path), WORKERS_SHEET)
    cols = _columns(table, WORKERS_COLUMNS, "trabajadores")
    result = ReadResult()
    for line, raw in enumerate(table[1:], start=2):

        def cell(title: str, raw=raw) -> object:
            return raw[cols[title]] if cols[title] < len(raw) else None

        if all(clean(v) == "" for v in raw):
            continue
        try:
            name = clean(cell("Nombre y apellidos"))
            if not name:
                raise ValueError("falta el nombre")
            area = clean(cell("Área"))
            if not area:
                raise ValueError("falta el área")
            result.rows.append(
                WorkerRow(
                    line,
                    name,
                    clean_id(cell("Carné de identidad")),
                    area,
                    clean(cell("Cargo")),
                    _group(cell("Grupo")),
                    _number(cell("Antigüedad"), "la antigüedad", 0, 60),
                    _number(cell("Días"), "los días", 0, 10_000),
                    _number(cell("Contribución"), "la contribución", 1, 5),
                    _yes_no(cell("Derecho a cobro")),
                )
            )
        except ValueError as error:
            result.errors.append(f"Fila {line}: {error}")
    return result


def read_structure(path: str | Path) -> ReadResult:
    """Rows of the «Áreas y cargos» table; rows that break the rules go to `errors`."""
    table = _load(Path(path), STRUCTURE_SHEET)
    cols = _columns(table, STRUCTURE_COLUMNS, "áreas y cargos")
    result = ReadResult()
    for line, raw in enumerate(table[1:], start=2):

        def cell(title: str, raw=raw) -> object:
            return raw[cols[title]] if cols[title] < len(raw) else None

        if all(clean(v) == "" for v in raw):
            continue
        try:
            area = clean(cell("Área"))
            if not area:
                raise ValueError("falta el área")
            row = StructureRow(line, area, clean(cell("Cargo")))
            row.approved = _number(cell("Aprobadas"), "las aprobadas", 0, 10_000, required=False) or 0
            row.vacant = _number(cell("Vacantes"), "las vacantes", 0, 10_000, required=False) or 0
            row.temp_vacant = _number(cell("Temporales"), "las temporales", 0, 10_000, required=False) or 0
            row.distribute = _number(
                cell("Vacantes a distribuir"), "las vacantes a distribuir", 0, 10_000, False
            )
            row.salary = _money(cell("Salario"))
            row.months_unexecuted = _number(cell("Meses"), "los meses", 0, 120, required=False) or 0
            total = row.vacant + row.temp_vacant
            if total > row.approved:
                raise ValueError(
                    f"las vacantes y temporales ({total}) superan las plazas aprobadas ({row.approved})"
                )
            if row.distribute is not None and row.distribute > total:
                raise ValueError(
                    f"las vacantes a distribuir ({row.distribute}) superan las vacantes ({total})"
                )
            result.rows.append(row)
        except ValueError as error:
            result.errors.append(f"Fila {line}: {error}")
    return result


# -- add-only plans ----------------------------------------------------------------------------
@dataclass
class WorkersPlan:
    new: list[WorkerRow] = field(default_factory=list)
    unchanged: int = 0  # already stored (same ID or same name): left as they are
    skipped: list[str] = field(default_factory=list)


@dataclass
class StructurePlan:
    new_areas: list[str] = field(default_factory=list)
    new_positions: list[StructureRow] = field(default_factory=list)
    unchanged: int = 0
    skipped: list[str] = field(default_factory=list)


def plan_workers(result: ReadResult, areas: list[Area], workers: list[Worker]) -> WorkersPlan:
    plan = WorkersPlan(skipped=list(result.errors))
    area_keys = {key(a.name) for a in areas}
    ids = {clean_id(w.national_id) for w in workers if w.national_id.strip()}
    id_by_name = {key(w.full_name): clean_id(w.national_id) for w in workers}
    for row in result.rows:
        known_id = id_by_name.get(key(row.name))
        # the ID identifies a person; the name only decides when one of the two has no ID
        same_name = known_id is not None and (
            not known_id or not row.national_id or known_id == row.national_id
        )
        if (row.national_id and row.national_id in ids) or same_name:
            plan.unchanged += 1
            continue
        if key(row.area) not in area_keys:
            plan.skipped.append(
                f"Fila {row.line}: el área «{row.area}» no existe; impórtela o créela primero en el paso 2."
            )
            continue
        plan.new.append(row)
        id_by_name[key(row.name)] = row.national_id
        if row.national_id:
            ids.add(row.national_id)
    return plan


def plan_structure(result: ReadResult, areas: list[Area], positions: list[Position]) -> StructurePlan:
    plan = StructurePlan(skipped=list(result.errors))
    area_by_key = {key(a.name): a for a in areas}
    known = {(p.area_id, key(p.name)) for p in positions}
    new_area_keys: set[str] = set()
    new_position_keys: set[tuple[str, str]] = set()
    for row in result.rows:
        area_key = key(row.area)
        existing = area_by_key.get(area_key)
        if existing is None and area_key not in new_area_keys:
            new_area_keys.add(area_key)
            plan.new_areas.append(row.area)
        if not row.position:
            plan.unchanged += int(existing is not None and area_key not in new_area_keys)
            continue
        position_key = key(row.position)
        if (existing is not None and (existing.id, position_key) in known) or (
            area_key,
            position_key,
        ) in new_position_keys:
            plan.unchanged += 1
            continue
        new_position_keys.add((area_key, position_key))
        plan.new_positions.append(row)
    return plan


# -- writing (report sheets and blank template) ------------------------------------------------
def _write_header(ws, titles: list[str], widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(titles)
    for index, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def add_data_sheets(wb, areas: list[Area], positions: list[Position], workers: list[Worker]) -> None:
    """«Trabajadores» and «Áreas y cargos» in the exact structure that the import reads."""
    from openpyxl.worksheet.datavalidation import DataValidation

    area_names = {a.id: a.name for a in areas}
    ws = wb.create_sheet(WORKERS_SHEET)
    _write_header(ws, WORKERS_COLUMNS, [34, 20, 34, 28, 9, 13, 8, 14, 16])
    for w in workers:
        group = to_roman(w.salary_group) if 1 <= w.salary_group <= 25 else ""
        ws.append(
            [
                w.full_name,
                w.national_id,
                area_names.get(w.area_id, ""),
                w.position_name,
                group,
                w.seniority_years,
                w.days_worked,
                w.contribution,
                YES if w.eligible else NO,
            ]
        )
    last = max(ws.max_row, 1000)
    for row in range(2, last + 1):  # ID numbers keep their leading zeros, also for rows typed later
        ws.cell(row=row, column=2).number_format = "@"
    rules = [
        ("E", DataValidation(type="list", formula1='"' + ",".join(to_roman(n) for n in range(1, 26)) + '"')),
        ("H", DataValidation(type="whole", operator="between", formula1="1", formula2="5")),
        ("I", DataValidation(type="list", formula1=f'"{YES},{NO}"')),
    ]
    for column, rule in rules:
        rule.allow_blank = True
        ws.add_data_validation(rule)
        rule.add(f"{column}2:{column}{last}")

    ws = wb.create_sheet(STRUCTURE_SHEET)
    _write_header(ws, STRUCTURE_COLUMNS, [34, 40, 11, 10, 12, 22, 12, 9])
    for area in areas:
        members = [p for p in positions if p.area_id == area.id]
        if not members:
            ws.append([area.name])
        for p in members:
            ws.append(
                [
                    area.name,
                    p.name,
                    p.approved,
                    p.vacant,
                    p.temp_vacant,
                    p.distribute,
                    float(p.salary),
                    p.months_unexecuted,
                ]
            )


INSTRUCTIONS = [
    "Cómo llenar este archivo para importarlo",
    "",
    "• No cambie los títulos de la primera fila ni el nombre de las hojas «Trabajadores» y «Áreas y cargos».",
    "• Primero importe «Áreas y cargos» (paso 2) y después «Trabajadores» (paso 3): "
    "cada trabajador debe estar en un área que ya exista.",
    "• La importación solo agrega lo que no existe. Lo que ya está en el programa "
    "(mismo carné o mismo nombre; "
    "misma área y cargo) no se cambia.",
    "",
    "Hoja «Trabajadores» (todas las columnas son obligatorias salvo Carné, Cargo y Derecho a cobro):",
    "  Grupo: número romano de I a XXV.   Antigüedad: años (entero).   Días: días trabajados en el período.",
    "  Contribución: 5 Excepcional, 4 Bueno, 3 Satisfactorio, 2 Regular, 1 Sin bonificación.",
    "  Derecho a cobro: Sí o No (si queda vacío, No).",
    "",
    "Hoja «Áreas y cargos» (una fila por cargo; una fila solo con el área crea el área sin cargos):",
    "  Aprobadas, Vacantes, Temporales, Meses: números enteros (vacío = 0).",
    "  Vacantes a distribuir: vacío = todas las vacantes.   Salario: monto mensual del cargo.",
]


def write_template(path: str | Path) -> None:
    """Blank workbook with the two sheets (titles only) and the instructions."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    add_data_sheets(wb, [], [], [])
    ws = wb.create_sheet(INSTRUCTIONS_SHEET)
    for line in INSTRUCTIONS:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    wb.save(path)
