"""Import workers from a user-provided CSV or XLSX file."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100}

FIELD_ALIASES = {
    "full_name": {
        "nombre",
        "nombreyapellidos",
        "nombrecompleto",
        "trabajador",
        "name",
        "fullname",
    },
    "national_id": {
        "ci",
        "carne",
        "carneidentidad",
        "carnedeidentidad",
        "identidad",
        "nationalid",
        "id",
    },
    "area": {"area", "departamento"},
    "position": {"cargo", "plaza", "puesto", "position"},
    "salary_group": {"grupo", "gruposalarial", "grupoescala", "salarygroup"},
    "seniority_years": {
        "antiguedad",
        "anos",
        "aniosantiguedad",
        "seniority",
        "seniorityyears",
    },
    "days_worked": {"dias", "diastrabajados", "tiempotrabajado", "daysworked"},
    "contribution": {"contribucion", "evaluacion", "contribution"},
}
REQUIRED = ("full_name", "area")


@dataclass
class WorkerRow:
    full_name: str
    national_id: str
    area: str
    position: str = ""
    salary_group: int | None = None
    seniority_years: int | None = None
    days_worked: int | None = None
    contribution: int | None = None
    area_number: int | None = None  # legacy sheet: number in the area header («11 Grupo Ecn.» -> 11)


@dataclass
class RowError:
    line: int
    message: str


@dataclass
class WorkerImportResult:
    rows: list[WorkerRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    legacy: bool = False  # read from the old per-worker sheet (areas come from its header rows)


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _map_headers(headers: list[object]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = _normalize_header(header)
        for field_name, aliases in FIELD_ALIASES.items():
            if key in aliases and field_name not in mapping:
                mapping[field_name] = index
    return mapping


def parse_roman(text: str) -> int | None:
    total, previous = 0, 0
    for char in reversed(text.upper()):
        value = ROMAN.get(char)
        if value is None:
            return None
        total += -value if value < previous else value
        previous = max(previous, value)
    return total


def _to_int(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        if float(value) != int(value):
            raise ValueError("not a whole number")
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+([.,]0+)?", text):
        return int(re.split(r"[.,]", text)[0])
    raise ValueError("not a whole number")


def _clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _normalize_national_id(value: object) -> str:
    if isinstance(value, float) and value == int(value):
        value = int(value)
    text = _clean_text(value)
    return text.zfill(11) if text.isdigit() and len(text) < 11 else text


def _parse_group(value: object) -> int | None:
    text = _clean_text(value)
    if not text:
        return None
    return parse_roman(text) if not text.isdigit() else int(text)


def _in_range(name: str, number: int | None, low: int, high: int) -> None:
    if number is not None and not low <= number <= high:
        raise ValueError(f"{name} must be between {low} and {high}")


def _build_row(raw: list[object], columns: dict[str, int]) -> WorkerRow:
    def cell(name: str) -> object:
        index = columns.get(name)
        return raw[index] if index is not None and index < len(raw) else None

    group = _parse_group(cell("salary_group"))
    if cell("salary_group") not in (None, "") and group is None:
        raise ValueError("salary group is not a valid number or roman numeral")
    _in_range("salary group", group, 1, 25)
    seniority = _to_int(cell("seniority_years"))
    _in_range("seniority", seniority, 0, 60)
    days = _to_int(cell("days_worked"))
    _in_range("days worked", days, 0, 10_000)
    contribution = _to_int(cell("contribution"))
    _in_range("contribution", contribution, 1, 5)
    return WorkerRow(
        full_name=_clean_text(cell("full_name")),
        national_id=_normalize_national_id(cell("national_id")),
        area=_clean_text(cell("area")),
        position=_clean_text(cell("position")),
        salary_group=group,
        seniority_years=seniority,
        days_worked=days,
        contribution=contribution,
    )


def _read_table(path: Path) -> list[list[object]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:  # CSV saved by Excel on Windows
        text = raw.decode("cp1252", errors="replace")
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text, newline=""), dialect)]


def read_workers(path: str | Path) -> WorkerImportResult:
    """Read workers; invalid rows are reported in `errors`, never raised."""
    table = _read_table(Path(path))
    result = WorkerImportResult()
    if not table:
        result.errors.append(RowError(1, "the file is empty"))
        return result
    columns = _map_headers(table[0])
    missing = [name for name in REQUIRED if name not in columns]
    if missing:
        result.errors.append(RowError(1, f"missing required columns: {', '.join(missing)}"))
        return result
    for line, raw in enumerate(table[1:], start=2):
        if all(_clean_text(value) == "" for value in raw):
            continue
        try:
            row = _build_row(raw, columns)
            if not row.full_name or not row.area:
                raise ValueError("name and area are required")
        except ValueError as error:
            result.errors.append(RowError(line, str(error)))
            continue
        result.rows.append(row)
    return result


# -- legacy "Dist.Monto x Trabajador" sheet ---------------------------------------------
NUMBERED_HEADER = re.compile(r"^\s*\d+[\s.,\-]")
AREA_NUMBER_PREFIX = re.compile(r"^\s*\d+\s*[.,\-]*\s*")


def normalize_area(name: object) -> str:
    """Comparable form of an area name: no «3.-» prefix, accents, case or punctuation."""
    return _normalize_header(AREA_NUMBER_PREFIX.sub("", _clean_text(name)))


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _find_legacy_sheet(workbook):
    for sheet in workbook.worksheets:
        key = _normalize_header(sheet.title)
        if "monto" in key and "trabajador" in key:
            return sheet
    return None


def read_legacy_workers(path: str | Path) -> WorkerImportResult | None:
    """Workers from the old per-worker sheet; None when the file has no such sheet.

    Area header rows have no days worked; worker rows carry the days in D:F (total in G), the roman
    salary group in I, the years in K and the contribution grade in M. Missing values stay None.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = _find_legacy_sheet(workbook)
        if sheet is None:
            return None
        result = WorkerImportResult(legacy=True)
        area, area_number = "", None
        for line, row in enumerate(sheet.iter_rows(min_col=1, max_col=14, values_only=True), start=1):
            name = _clean_text(row[2])
            if not name:
                continue
            days_parts = [v for v in row[3:6] if _is_number(v)]
            total = row[6] if _is_number(row[6]) and row[6] > 0 else sum(days_parts)
            if total <= 0:
                # only numbered headers start a new area; unnumbered ones are sub-groups of the current one
                if "monto" not in name.lower() and NUMBERED_HEADER.match(name):
                    area = AREA_NUMBER_PREFIX.sub("", name).strip()
                    area_number = int(re.match(r"\s*(\d+)", name).group(1))
                continue
            if name.startswith("#"):
                result.errors.append(RowError(line, "the worker name is a broken reference (#REF!)"))
                continue
            group = parse_roman(_clean_text(row[8])) if _clean_text(row[8]) else None
            years = int(row[10]) if _is_number(row[10]) and row[10] >= 0 else None
            contribution = int(row[12]) if _is_number(row[12]) and 1 <= row[12] <= 5 else None
            result.rows.append(
                WorkerRow(
                    full_name=name,
                    national_id="",
                    area=area,
                    salary_group=group if group and 1 <= group <= 25 else None,
                    seniority_years=years,
                    days_worked=int(total),
                    contribution=contribution,
                    area_number=area_number,
                )
            )
        return result
    finally:
        workbook.close()


def _find_roster_header(sheet):
    """(row index, {field: column}) of a «Areas y Cargos | Trabajador | Carnet | Grupo» header, if any."""
    for index, row in enumerate(sheet.iter_rows(max_row=15, values_only=True)):
        keys = [_normalize_header(c) if c is not None else "" for c in row]
        columns: dict[str, int] = {}
        for col, key in enumerate(keys):
            if "area" in key and "cargo" in key:
                columns["position"] = col
            elif key.startswith("trabajador"):
                columns["name"] = col
            elif "carnet" in key or "carne" in key:
                columns["id"] = col
            elif key.startswith("grupo"):
                columns["group"] = col
        if {"position", "name"} <= columns.keys():
            return index, columns
    return None


def read_roster_workers(path: str | Path) -> WorkerImportResult | None:
    """Workers from a staffing roster: area rows only have the first column, position rows also have the
    salary group, and a position row with a worker name is filled. None when no sheet looks like this."""
    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            found = _find_roster_header(sheet)
            if found is None:
                continue
            header, cols = found
            result = WorkerImportResult(legacy=True)
            area = ""
            for line, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if line <= header + 1:
                    continue

                def cell(key: str, row=row, cols=cols) -> object:
                    col = cols.get(key)
                    return row[col] if col is not None and col < len(row) else None

                title = _clean_text(cell("position"))
                if not title:
                    continue
                group = parse_roman(_clean_text(cell("group"))) if _clean_text(cell("group")) else None
                if not _clean_text(cell("group")) and not _clean_text(cell("name")):
                    area = title  # a row with only the first column starts a new area
                    continue
                name = _clean_text(cell("name"))
                if not name:
                    continue  # vacant position
                result.rows.append(
                    WorkerRow(
                        full_name=name,
                        national_id=_normalize_national_id(cell("id")),
                        area=area,
                        position=title,
                        salary_group=group if group and 1 <= group <= 25 else None,
                    )
                )
            return result
        return None
    finally:
        workbook.close()


def read_workers_auto(path: str | Path) -> WorkerImportResult:
    """Legacy per-worker sheet or staffing roster when the workbook has one, otherwise a plain table."""
    if Path(path).suffix.lower() in {".xlsx", ".xlsm"}:
        for reader in (read_legacy_workers, read_roster_workers):
            found = reader(path)
            if found is not None:
                return found
    return read_workers(path)
