"""Window-level tests. They need a display (Xvfb on Linux); they are skipped when tkinter cannot start."""

import tkinter as tk
from decimal import Decimal as D

import pytest

from fondo_salario.domain.models import Position, Settings, Worker
from fondo_salario.persistence.sqlite_repo import ProjectStore
from fondo_salario.ui import app as app_module


@pytest.fixture
def window(tmp_path, monkeypatch):
    try:
        probe = tk.Tk()
        probe.destroy()
    except tk.TclError:
        pytest.skip("no display available")
    path = tmp_path / "p.fsal"
    with ProjectStore(path) as store:
        store.add_area("Dir")
    win = app_module.App(path)
    win.withdraw()
    win.update()
    yield win
    try:
        win.destroy()
    except tk.TclError:
        pass


def ask(monkeypatch, answer):
    monkeypatch.setattr(app_module.messagebox, "askyesnocancel", lambda *a, **k: answer)


def reopen(path):
    with ProjectStore(path) as store:
        return store.load()


def test_opens_with_four_steps_and_visible_file_controls(window):
    assert list(window.tabs) == ["1. Parámetros", "2. Áreas y cargos", "3. Trabajadores", "4. Resultados"]
    assert window.status.cget("text").endswith("Todo guardado")


def test_closing_with_changes_asks_and_cancel_keeps_window_open(window, monkeypatch):
    window.ctl.add_area("Nueva")
    window.update_status()
    assert "sin guardar" in window.status.cget("text")
    ask(monkeypatch, None)  # Cancelar
    window.on_close()
    assert window.winfo_exists()


def test_closing_and_saving_persists_changes(window, monkeypatch):
    window.ctl.add_area("Nueva")
    ask(monkeypatch, True)  # Sí
    path = window.project_path
    window.on_close()
    assert [a.name for a in reopen(path).areas] == ["Dir", "Nueva"]


def test_closing_without_saving_discards_changes(window, monkeypatch):
    window.ctl.add_area("Nueva")
    ask(monkeypatch, False)  # No
    path = window.project_path
    window.on_close()
    assert [a.name for a in reopen(path).areas] == ["Dir"]


def test_closing_with_nothing_pending_does_not_ask(window, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("should not ask")

    monkeypatch.setattr(app_module.messagebox, "askyesnocancel", boom)
    window.on_close()


def test_edited_parameters_are_applied_on_save(window):
    params = window.tabs["1. Parámetros"]
    params.vars["total"].set("500.25")
    assert params.edited and "sin guardar" in window.status.cget("text")
    assert window.save() is True
    assert window.ctl.project.settings.total_amount == D("500.25") and not params.edited
    assert reopen(window.project_path).settings.total_amount == D("500.25")


def test_invalid_parameters_block_saving(window, monkeypatch):
    errors = []
    monkeypatch.setattr(app_module.messagebox, "showerror", lambda *a, **k: errors.append(a))
    window.tabs["1. Parámetros"].vars["months"].set("abc")
    assert window.save() is False and errors


def test_results_tab_balances_with_vacancies_to_distribute(window):
    window.ctl.save_settings(Settings(total_amount=D("1000.00")))
    area = window.ctl.project.areas[0]
    window.ctl.add_position(Position(0, area.id, "Jefe", 2, 3, 1, D("100"), 3, distribute=2))
    window.ctl.add_worker(Worker(0, "Ana Prueba", "1", area.id, 12, 5, 72, 4))
    window.changed()
    results = window.tabs["4. Resultados"]
    assert "Cuadra" in results.summary.cget("text") and "No cuadra" not in results.summary.cget("text")
    assert (
        window.tabs["2. Áreas y cargos"].tree.item(f"p{window.ctl.project.positions[0].id}", "values")[3]
        == "2"
    )


def test_selftest_opens_and_closes_the_window(tmp_path):
    try:
        probe = tk.Tk()
        probe.destroy()
    except tk.TclError:
        pytest.skip("no display available")
    assert app_module.selftest() == 0
    assert app_module.main(["--selftest"]) == 0


def test_results_do_not_claim_balance_when_the_total_is_missing(window):
    area = window.ctl.project.areas[0]
    window.ctl.add_position(Position(0, area.id, "Jefe", 1, 1, 0, D("100"), 3))
    window.ctl.add_worker(Worker(0, "Ana Prueba", "1", area.id, 12, 5, 72, 4))
    window.changed()
    text = window.tabs["4. Resultados"].summary.cget("text")
    assert "Falta" in text and "Cuadra" not in text


def test_reset_button_restores_the_regulation_values(window):
    params = window.tabs["1. Parámetros"]
    params.vars["w_time"].set("0.50")
    params.vars["b_sen"].set("0:1, 10:5")
    params.use_regulation_values()
    assert params.vars["w_time"].get() == "0.20" and params.vars["w_con"].get() == "0.40"
    assert params.vars["b_sen"].get() == "0:1, 6:2, 11:3, 16:4, 21:5"


def _add_workers(window, count=3):
    area = window.ctl.project.areas[0]
    for i in range(count):
        window.ctl.add_worker(Worker(0, f"Trabajador {i}", str(i), area.id, 12, 5, 72, 4))
    window.changed()
    return window.tabs["3. Trabajadores"]


def test_bulk_toggle_of_the_right_to_collect(window):
    tab = _add_workers(window)
    ids = tab.tree.get_children()
    tab.tree.selection_set(ids[:2])
    tab.revoke_right()
    assert [w.eligible for w in window.ctl.project.workers] == [False, False, True]
    tab.tree.selection_set(ids[:1])
    tab.grant_right()
    assert [w.eligible for w in window.ctl.project.workers] == [True, False, True]


def test_filter_by_right_to_collect(window):
    tab = _add_workers(window)
    window.ctl.set_eligible([window.ctl.project.workers[0].id], False)
    window.changed()
    tab.right_filter.set("Con derecho")
    tab.refresh()
    assert len(tab.tree.get_children()) == 2
    tab.right_filter.set("Sin derecho")
    tab.refresh()
    assert len(tab.tree.get_children()) == 1


def test_contribution_labels_follow_the_regulation():
    assert app_module.CONTRIBUTION_CHOICES == [
        "5 - Excepcional",
        "4 - Bueno",
        "3 - Satisfactorio",
        "2 - Regular",
        "1 - Sin bonificación",
    ]
    assert app_module.contribution_from_text("4 - Bueno") == 4
    assert app_module.contribution_from_text("1 - Sin bonificación") == 1
    assert app_module.contribution_label(3) == "3 - Satisfactorio"


# --- import with the fixed structure -------------------------------------------------------
def _book(tmp_path, workers=(), structure=()):
    from openpyxl import Workbook

    from fondo_salario.importers.table import (
        STRUCTURE_COLUMNS,
        STRUCTURE_SHEET,
        WORKERS_COLUMNS,
        WORKERS_SHEET,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = WORKERS_SHEET
    ws.append(WORKERS_COLUMNS)
    for row in workers:
        ws.append(row)
    ws = wb.create_sheet(STRUCTURE_SHEET)
    ws.append(STRUCTURE_COLUMNS)
    for row in structure:
        ws.append(row)
    path = tmp_path / "datos.xlsx"
    wb.save(path)
    return path


def _patch_dialogs(monkeypatch, path, agree=True):
    monkeypatch.setattr(app_module.filedialog, "askopenfilename", lambda *a, **k: str(path))
    seen = {"info": [], "error": [], "question": []}
    monkeypatch.setattr(app_module.messagebox, "showinfo", lambda t, m, **k: seen["info"].append(m))
    monkeypatch.setattr(app_module.messagebox, "showerror", lambda t, m, **k: seen["error"].append(m))

    def ask_ok(title, message, **kwargs):
        seen["question"].append(message)
        return agree

    monkeypatch.setattr(app_module.messagebox, "askokcancel", ask_ok)
    return seen


WORKER_ROW = ["Ana Uno", "00112233445", "Dir", "Técnico", "XII", 5, 72, 3, "Sí"]


def test_workers_import_adds_new_ones_after_confirming(window, monkeypatch, tmp_path):
    path = _book(tmp_path, workers=[WORKER_ROW, ["Beto Dos", "", "Dir", "", "III", 1, 72, 4, "No"]])
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["3. Trabajadores"].import_file()
    by_name = {w.full_name: w for w in window.ctl.project.workers}
    assert set(by_name) == {"Ana Uno", "Beto Dos"}
    assert (by_name["Ana Uno"].eligible, by_name["Beto Dos"].eligible) == (True, False)
    assert by_name["Ana Uno"].salary_group == 12 and "Trabajadores nuevos: 2" in seen["question"][0]
    assert "Trabajadores agregados: 2" in seen["info"][-1]


def test_import_never_changes_a_worker_that_already_exists(window, monkeypatch, tmp_path):
    area = window.ctl.project.areas[0]
    window.ctl.add_worker(Worker(0, "Ana Uno", "", area.id, 20, 9, 10, 5))
    path = _book(tmp_path, workers=[WORKER_ROW])
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["3. Trabajadores"].import_file()
    (ana,) = window.ctl.project.workers
    assert (ana.salary_group, ana.seniority_years, ana.national_id) == (20, 9, "")
    assert seen["question"] == [] and "No hay registros nuevos" in seen["info"][-1]


def test_declining_the_confirmation_adds_nothing(window, monkeypatch, tmp_path):
    path = _book(tmp_path, workers=[WORKER_ROW])
    _patch_dialogs(monkeypatch, path, agree=False)
    window.tabs["3. Trabajadores"].import_file()
    assert window.ctl.project.workers == []


def test_a_file_that_breaks_the_structure_is_refused_with_the_missing_columns(window, monkeypatch, tmp_path):
    path = tmp_path / "otro.csv"
    path.write_text("nombre;area;edad\nAna;Dir;3\n", encoding="utf-8")
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["3. Trabajadores"].import_file()
    assert window.ctl.project.workers == []
    assert "Faltan estas columnas" in seen["error"][0] and "Carné de identidad" in seen["error"][0]


def test_rows_with_problems_are_listed_and_the_rest_is_imported(window, monkeypatch, tmp_path):
    bad_area = ["Carla Tres", "", "No existe", "", "III", 1, 72, 4, "No"]
    bad_group = ["Dani Cuatro", "", "Dir", "", "ZZZ", 1, 72, 4, "No"]
    path = _book(tmp_path, workers=[WORKER_ROW, bad_area, bad_group])
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["3. Trabajadores"].import_file()
    assert [w.full_name for w in window.ctl.project.workers] == ["Ana Uno"]
    assert "el área «No existe» no existe" in seen["question"][0] and "grupo" in seen["question"][0]


def test_structure_import_adds_areas_and_positions(window, monkeypatch, tmp_path):
    structure = [
        ["Dir", "Jefe", 2, 1, 0, None, 6060, 3],
        ["Economía", "Contador", 1, 1, 0, 1, 4000, 3],
        ["Archivo"],
    ]
    path = _book(tmp_path, structure=structure)
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["2. Áreas y cargos"].import_excel()
    project = window.ctl.project
    assert [a.name for a in project.areas] == ["Dir", "Economía", "Archivo"]
    assert {(p.name, p.vacant, p.salary) for p in project.positions} == {
        ("Jefe", 1, D("6060")),
        ("Contador", 1, D("4000")),
    }
    assert "Áreas nuevas: 2" in seen["question"][0] and "Cargos nuevos: 2" in seen["question"][0]
    assert window.ctl.project.positions[0].distribute is None


def test_template_has_the_structure_and_can_be_saved(window, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "plantilla.xlsx"
    monkeypatch.setattr(app_module.filedialog, "asksaveasfilename", lambda *a, **k: str(out))
    monkeypatch.setattr(app_module.messagebox, "showinfo", lambda *a, **k: None)
    window.save_template()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Trabajadores", "Áreas y cargos", "Instrucciones"]
    assert [c.value for c in wb["Trabajadores"][1]][:3] == [
        "Nombre y apellidos",
        "Carné de identidad",
        "Área",
    ]


def test_clear_workers_asks_first_and_keeps_areas(window, monkeypatch):
    area = window.ctl.project.areas[0]
    window.ctl.add_worker(Worker(0, "Ana Uno", "", area.id, 12, 5, 72, 3))
    tab = window.tabs["3. Trabajadores"]
    monkeypatch.setattr(app_module.messagebox, "askyesno", lambda *a, **k: False)
    tab.clear_all()
    assert len(window.ctl.project.workers) == 1  # declined
    monkeypatch.setattr(app_module.messagebox, "askyesno", lambda *a, **k: True)
    tab.clear_all()
    assert window.ctl.project.workers == [] and len(window.ctl.project.areas) == 1


def test_clear_everything_removes_areas_positions_and_workers(window, monkeypatch):
    area = window.ctl.project.areas[0]
    window.ctl.add_position(Position(0, area.id, "Jefe", 1, 1, 0, D("100"), 1))
    window.ctl.add_worker(Worker(0, "Ana Uno", "", area.id, 12, 5, 72, 3))
    asked = []
    monkeypatch.setattr(app_module.messagebox, "askyesno", lambda t, m, **k: asked.append(m) or True)
    window.tabs["2. Áreas y cargos"].clear_all()
    project = window.ctl.project
    assert project.areas == [] and project.positions == [] and project.workers == []
    assert "TODAS las áreas" in asked[0] and "trabajadores (1)" in asked[0]
    monkeypatch.setattr(app_module.messagebox, "showinfo", lambda *a, **k: asked.append("info"))
    window.tabs["2. Áreas y cargos"].clear_all()  # nothing left: says so instead of asking
    assert asked[-1] == "info"


def test_structure_file_without_the_titles_is_refused(window, monkeypatch, tmp_path):
    path = tmp_path / "areas.csv"
    path.write_text("area;puesto\nDir;Jefe\n", encoding="utf-8")
    seen = _patch_dialogs(monkeypatch, path)
    window.tabs["2. Áreas y cargos"].import_excel()
    assert "Faltan estas columnas" in seen["error"][0] and "Vacantes a distribuir" in seen["error"][0]
    assert window.ctl.project.positions == []


def test_a_failing_write_is_reported_and_the_screen_refreshes(window, monkeypatch, tmp_path):
    path = _book(tmp_path, workers=[WORKER_ROW])
    seen = _patch_dialogs(monkeypatch, path)

    def boom(plan):
        raise RuntimeError("disco lleno")

    monkeypatch.setattr(window.ctl, "apply_workers_import", boom)
    window.tabs["3. Trabajadores"].import_file()
    assert "interrumpió" in seen["error"][0] and "disco lleno" in seen["error"][0]
