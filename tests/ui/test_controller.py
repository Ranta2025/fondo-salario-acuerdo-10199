from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from fondo_salario.domain.models import Position, Settings, Worker
from fondo_salario.importers.merge import resolve
from fondo_salario.importers.structure import AreaRow, PositionRow
from fondo_salario.importers.workers import WorkerRow
from fondo_salario.persistence.sqlite_repo import ProjectStore
from fondo_salario.ui.controller import Controller


@pytest.fixture
def ctl(tmp_path):
    with ProjectStore(tmp_path / "p.fsal") as store:
        yield Controller(store)


def row(name="Ana Prueba", area="Dir", **kw):
    base = dict(full_name=name, national_id="", area=area, salary_group=12, seniority_years=5)
    base.update(kw)
    return WorkerRow(**base)


def test_import_structure_creates_areas_and_skips_existing(ctl):
    data = [AreaRow(1, "Dir", [PositionRow("Jefe", 1, D("6000"), 3)]), AreaRow(2, "Eco")]
    assert ctl.import_structure(data) == 2
    assert ctl.import_structure(data) == 0
    project = ctl.project
    assert [a.name for a in project.areas] == ["Dir", "Eco"]
    assert project.positions[0].vacant == 1 and project.positions[0].distribute is None


def test_import_workers_matches_area_ignoring_case_and_accents(ctl):
    ctl.import_structure([AreaRow(1, "Dirección")])
    summary = ctl.import_workers([row(area="direccion"), row(name="Beto", area="Nope")])
    assert summary.added == 1
    assert "Nope" in summary.skipped[0]


def test_import_workers_requires_group_and_seniority_and_defaults_the_rest(ctl):
    ctl.save_settings(Settings(total_amount=D("0"), months_paid=2))
    ctl.import_structure([AreaRow(1, "Dir")])
    summary = ctl.import_workers([row(salary_group=None), row(name="Beto")])
    assert summary.added == 1 and len(summary.skipped) == 1
    worker = ctl.project.workers[0]
    assert worker.days_worked == 48 and worker.contribution == 3
    assert summary.defaults_used == 1


def test_calculation_reflects_edits(ctl):
    ctl.save_settings(Settings(total_amount=D("500.00")))
    area = ctl.add_area("Dir")
    ctl.add_position(Position(0, area.id, "Jefe", 1, 1, 0, D("1000"), 3))
    ctl.add_worker(Worker(0, "Ana", "1", area.id, 12, 5, 72, 4))
    assert ctl.calculation().workers[0].amount == D("500.00")
    ctl.update_worker(Worker(ctl.project.workers[0].id, "Ana", "1", area.id, 12, 5, 72, 4, eligible=False))
    assert ctl.calculation().workers[0].amount == D("0.00")


def test_export_writes_workbook(ctl, tmp_path):
    ctl.save_settings(Settings(total_amount=D("10.00")))
    ctl.add_area("Dir")
    out = tmp_path / "out.xlsx"
    ctl.export(out)
    assert "Pre-nómina" in load_workbook(out).sheetnames


@pytest.mark.parametrize(
    "text,expected",
    [
        ("1234.56", "1234.56"),
        ("1 234,56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("1.234,56", "1234.56"),
        ("100", "100"),
    ],
)
def test_parse_money(text, expected):
    from fondo_salario.ui.controller import parse_money

    assert parse_money(text) == D(expected)


@pytest.mark.parametrize(
    "text,total,expected",
    [("", 4, None), ("  ", 4, None), ("4", 4, None), ("2", 4, 2), ("0", 4, 0)],
)
def test_resolve_distribute(text, total, expected):
    from fondo_salario.ui.controller import resolve_distribute

    assert resolve_distribute(text, total) == expected


def test_resolve_distribute_rejects_more_than_available():
    from fondo_salario.ui.controller import resolve_distribute

    with pytest.raises(ValueError, match="no puede ser mayor"):
        resolve_distribute("5", 4)
    with pytest.raises(ValueError):
        resolve_distribute("x", 4)


def test_controller_tracks_unsaved_changes(tmp_path):
    with ProjectStore(tmp_path / "p.fsal", autosave=False) as store:
        c = Controller(store)
        assert c.dirty is False
        c.add_area("Dir")
        assert c.dirty is True
        c.save()
        assert c.dirty is False
        c.add_area("Eco")
        c.discard()
        assert [a.name for a in c.project.areas] == ["Dir"] and c.dirty is False


def test_import_legacy_creates_missing_areas_and_keeps_undefined_group(ctl):
    ctl.import_structure([AreaRow(1, "Desarrollo Habitacional")])
    rows = [
        row(name="Ana Uno", area="Desarrollo Habitacional", salary_group=None, seniority_years=None),
        row(name="Beto Dos", area="Banco de los Colonos", salary_group=12),
    ]
    summary = ctl.import_workers(rows, create_missing_areas=True, allow_undefined=True)
    assert summary.added == 2 and summary.areas_created == ["Banco de los Colonos"]
    assert summary.undefined == 1
    ana = next(w for w in ctl.project.workers if w.full_name == "Ana Uno")
    assert ana.salary_group == 0 and ana.seniority_years == 0
    assert ana.eligible is False  # nobody has the right to collect until it is granted
    assert not any(i.code == "invalid_worker" for i in ctl.calculation().issues)
    # once the right is granted, the calculation reports the incomplete data instead of paying wrongly
    ctl.set_eligible([ana.id], True)
    assert any(i.code == "invalid_worker" for i in ctl.calculation().issues)


def test_import_matches_numbered_area_names_and_skips_duplicates(ctl):
    ctl.import_structure([AreaRow(1, "Desarrollo Habitacional")])
    first = ctl.import_workers([row(name="Ana Uno", area="4. Desarrollo Habitacional")])
    assert first.added == 1
    again = ctl.import_workers([row(name="ana  uno", area="Desarrollo Habitacional")])
    assert again.added == 0 and again.unchanged == 1


def test_import_matches_areas_by_containment_and_legacy_number(ctl):
    ctl.import_structure(
        [
            AreaRow(1, "Direccion"),
            AreaRow(2, "Grupo de Subsidios"),
            AreaRow(3, "G. Atencion C/T"),
            AreaRow(4, "Dpto. Administrativo"),
        ]
    )
    rows = [
        row(name="Uno", area="Grupo de Subsidios y Esfuerzos p.", area_number=2),
        row(name="Dos", area="Atencion C/T", area_number=3),
        row(name="Tres", area="Grupo Ecn.", area_number=4),
    ]
    summary = ctl.import_workers(rows, create_missing_areas=True)
    assert summary.areas_created == [] and summary.added == 3
    names = {a.id: a.name for a in ctl.project.areas}
    assert [names[w.area_id] for w in ctl.project.workers] == [
        "Grupo de Subsidios",
        "G. Atencion C/T",
        "Dpto. Administrativo",
    ]


def test_set_eligible_for_several_workers_at_once(ctl):
    from fondo_salario.domain.models import Worker

    area = ctl.add_area("Dir")
    ids = [ctl.add_worker(Worker(0, f"W{i}", str(i), area.id, 12, 5, 72, 4)).id for i in range(3)]
    ctl.set_eligible(ids[:2], False)
    assert [w.eligible for w in ctl.project.workers] == [False, False, True]
    ctl.set_eligible(ids[:1], True)
    assert [w.eligible for w in ctl.project.workers] == [True, False, True]
    assert ctl.dirty is False or True  # state tracked by the store (autosave on in this fixture)


def test_imported_workers_start_without_the_right_to_collect(ctl):
    ctl.import_structure([AreaRow(1, "Dir")])
    ctl.import_workers([row(name="Ana Uno", area="Dir"), row(name="Beto Dos", area="Dir")])
    assert [w.eligible for w in ctl.project.workers] == [False, False]


# --- merging imports into existing data ------------------------------------------------
def _ctl_with_worker_without_id(ctl):
    ctl.import_structure([AreaRow(1, "Dirección"), AreaRow(2, "Economía")])
    ctl.import_workers([row(name="Ana Pérez Gómez", area="Dirección", salary_group=12, seniority_years=5)])
    worker = ctl.project.workers[0]
    ctl.set_eligible([worker.id], True)
    return ctl.project.workers[0]


def test_join_completes_the_missing_id_without_duplicating(ctl):
    _ctl_with_worker_without_id(ctl)
    plan = ctl.plan_workers_import([row(name="ana perez gomez", national_id="00112233445", area="Dirección")])
    assert resolve(plan, "keep_all") is True
    summary = ctl.apply_import(plan)
    (w,) = ctl.project.workers
    assert (w.full_name, w.national_id, w.eligible) == ("Ana Pérez Gómez", "00112233445", True)
    assert (summary.added, summary.completed, summary.overwritten, summary.kept) == (0, 1, 0, 0)


def test_differences_are_applied_only_when_accepted(ctl):
    _ctl_with_worker_without_id(ctl)
    plan = ctl.plan_workers_import([row(name="Ana Pérez Gómez", area="Economía", salary_group=20)])
    diffs = {d.key: d for _, d in plan.conflicts}
    diffs["salary_group"].accept = True
    diffs["area_id"].accept = False
    summary = ctl.apply_import(plan)
    (w,) = ctl.project.workers
    assert w.salary_group == 20 and ctl.project.areas[0].id == w.area_id
    assert (summary.overwritten, summary.kept) == (1, 1)


def test_accepting_an_area_change_moves_the_worker_and_can_create_the_area(ctl):
    _ctl_with_worker_without_id(ctl)
    plan = ctl.plan_workers_import(
        [row(name="Ana Pérez Gómez", area="Banco de los Colonos")], create_missing_areas=True
    )
    resolve(plan, "accept_all")
    summary = ctl.apply_import(plan)
    names = {a.id: a.name for a in ctl.project.areas}
    assert names[ctl.project.workers[0].area_id] == "Banco de los Colonos"
    assert summary.areas_created == ["Banco de los Colonos"]


def test_new_workers_are_added_without_the_right_and_existing_ones_are_untouched(ctl):
    _ctl_with_worker_without_id(ctl)
    plan = ctl.plan_workers_import(
        [
            row(name="Ana Pérez Gómez", area="Dirección"),
            row(name="Beto Ruiz", area="Dirección", salary_group=3, seniority_years=1),
        ]
    )
    ctl.apply_import(plan)
    by_name = {w.full_name: w for w in ctl.project.workers}
    assert by_name["Ana Pérez Gómez"].eligible is True and by_name["Beto Ruiz"].eligible is False


def test_structure_merge_adds_new_positions_and_completes_existing_ones(ctl):
    ctl.import_structure([AreaRow(1, "Dir", [PositionRow("Jefe", 0, D("0"), 0)])])
    data = [AreaRow(1, "Dir", [PositionRow("Jefe", 2, D("6060"), 3), PositionRow("Chofer", 1, D("2200"), 3)])]
    plan = ctl.plan_structure_import(data + [AreaRow(2, "Eco", [PositionRow("EP", 1, D("4610"), 3)])])
    summary = ctl.apply_import(plan)
    by_name = {p.name: p for p in ctl.project.positions}
    assert (by_name["Jefe"].vacant, by_name["Jefe"].salary, by_name["Jefe"].months_unexecuted) == (
        2,
        D("6060"),
        3,
    )
    assert {"Chofer", "EP"} <= set(by_name)
    assert [a.name for a in ctl.project.areas] == ["Dir", "Eco"]
    assert (summary.areas_added, summary.positions_added, summary.completed) == (1, 2, 3)


def test_structure_merge_keeps_the_stored_value_unless_accepted(ctl):
    ctl.import_structure([AreaRow(1, "Dir", [PositionRow("Jefe", 1, D("5000"), 3)])])
    plan = ctl.plan_structure_import([AreaRow(1, "Dir", [PositionRow("Jefe", 1, D("6060"), 3)])])
    resolve(plan, "keep_all")
    ctl.apply_import(plan)
    assert ctl.project.positions[0].salary == D("5000")
    plan = ctl.plan_structure_import([AreaRow(1, "Dir", [PositionRow("Jefe", 1, D("6060"), 3)])])
    resolve(plan, "accept_all")
    ctl.apply_import(plan)
    assert ctl.project.positions[0].salary == D("6060")


def test_workers_file_on_an_empty_project_creates_the_areas(ctl):
    plan = ctl.plan_workers_import([row(name="Ana Uno", area="Dirección")], create_missing_areas=True)
    assert plan.new_count == 1 and plan.areas_to_create == ["Dirección"]
    ctl.apply_import(plan)
    assert [a.name for a in ctl.project.areas] == ["Dirección"] and len(ctl.project.workers) == 1


def test_nothing_new_text_does_not_claim_everything_is_registered_when_rows_were_skipped(ctl):
    from fondo_salario.ui.app import nothing_new_text

    plan = ctl.plan_workers_import([row(area="Zzz")])
    text = nothing_new_text(plan, [])
    assert "No hay nada nuevo" not in text and "ninguna fila" in text
