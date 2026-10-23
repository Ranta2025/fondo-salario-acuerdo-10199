from decimal import Decimal

from openpyxl import Workbook

from fondo_salario.importers.structure import read_structure
from fondo_salario.importers.workers import parse_roman, read_workers


def test_parse_roman():
    assert [parse_roman(x) for x in ("IV", "XII", "XIX", "XXV", "abc")] == [
        4,
        12,
        19,
        25,
        None,
    ]


def test_csv_with_accents_aliases_and_semicolon(tmp_path):
    path = tmp_path / "w.csv"
    path.write_text(
        "Nombre y Apellidos;Carné;Área;Cargo;Grupo;Antigüedad;Días;Contribución\n"
        "Ana Prueba Uno;123;Dir;Jefe;XII;5;72;4\n"
        ";;;;;;;\n",
        encoding="utf-8",
    )
    result = read_workers(path)
    assert result.errors == []
    row = result.rows[0]
    assert (row.full_name, row.national_id, row.area, row.salary_group) == (
        "Ana Prueba Uno",
        "00000000123",
        "Dir",
        12,
    )
    assert (row.seniority_years, row.days_worked, row.contribution) == (5, 72, 4)


def test_invalid_rows_are_reported_not_raised(tmp_path):
    path = tmp_path / "w.csv"
    path.write_text("nombre,area,contribucion\nA,X,9\n,X,3\nB,X,3\n", encoding="utf-8")
    result = read_workers(path)
    assert [r.full_name for r in result.rows] == ["B"]
    assert [e.line for e in result.errors] == [2, 3]


def test_missing_required_columns(tmp_path):
    path = tmp_path / "w.csv"
    path.write_text("foo,bar\n1,2\n", encoding="utf-8")
    assert "missing required columns" in read_workers(path).errors[0].message


def test_xlsx_numeric_id_keeps_leading_zeros(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre", "CI", "Area"])
    ws.append(["Luis Prueba", 1234567890.0, "Dir"])
    path = tmp_path / "w.xlsx"
    wb.save(path)
    assert read_workers(path).rows[0].national_id == "01234567890"


def test_legacy_structure_layout(tmp_path):
    wb = Workbook()
    wb.active.title = "Otra"
    ws = wb.create_sheet("Dist.Monto x area")
    ws["B3"], ws["C3"] = None, "Areas"
    ws.append([])
    rows = [
        (1, "1.-Direccion", None, None, None, None),
        (1, "SubDirector", 1, 6310, None, 3),
        (1, "Secretaria", None, None, None, None),
        (2, "2.- Grupo de Atencion", None, None, None, None),
        (2, "Tecnico Integral", 4, 3610, None, 3),
    ]
    for number, name, approved, salary, extra, months in rows:
        ws.append([None, number, name, approved, salary, extra, months])
    ws.append([None, None, "Total", 5, 9920])
    path = tmp_path / "s.xlsx"
    wb.save(path)

    areas = read_structure(path)
    assert [a.name for a in areas] == ["Direccion", "Grupo de Atencion"]
    sub = areas[0].positions[0]
    assert (sub.name, sub.vacant, sub.salary, sub.months_unexecuted) == (
        "SubDirector",
        1,
        Decimal(6310),
        3,
    )
    assert [p.name for p in areas[0].positions] == ["SubDirector", "Secretaria"]
    assert areas[1].positions[0].vacant == 4


def test_csv_saved_by_excel_on_windows_cp1252(tmp_path):
    path = tmp_path / "w.csv"
    path.write_bytes("Nombre;Área;Antigüedad\nMaría Núñez Prueba;Dirección;5\n".encode("cp1252"))
    result = read_workers(path)
    assert result.errors == []
    assert (result.rows[0].full_name, result.rows[0].area) == ("María Núñez Prueba", "Dirección")


# --- legacy "Dist.Monto x Trabajador" sheet ------------------------------------------
def _legacy_workers_workbook(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Relacion de Areas"
    ws = wb.create_sheet("Dist.Monto x Trabajador")
    ws["D2"], ws["I2"] = "1.- Tiempo Trabajado", "2.- Responsabilidad"
    ws["C4"], ws["O4"] = "Monto Total a Distribuir", 1000
    #       A  B  C                                D   E   F   G   H  I      J  K   L  M
    ws.append([])
    rows = [
        ("x", None, "3.- Departamento de Control", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        (None, None, "Nidia Prueba Uno", 24, 24, 24, 72, 5, "XII", 3, 21, 5, 3),
        (None, None, None, None, None, None, None, None, None, "Coeficiente del Grupo", None, None, None),
        (
            None,
            None,
            "4. Desarrollo Habitacional",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ),
        (None, None, "Sin Grupo Prueba", 24, 24, 24, 72, 5, None, 4, None, 4, 9),
        (None, None, "#REF!", 24, 24, 24, 72, 5, None, 4, 3, 4, 4),
        (None, None, "Banco de los Colonos", None, None, None, None, None, None, None, None, None, None),
        (None, None, "Maria Prueba Dos ", 24, 24, 12, None, 5, "XX", 4, 0, 5, 4),
    ]
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "legacy.xlsx"
    wb.save(path)
    return path


def test_legacy_workers_sheet(tmp_path):
    from fondo_salario.importers.workers import read_legacy_workers

    result = read_legacy_workers(_legacy_workers_workbook(tmp_path))
    by_name = {r.full_name: r for r in result.rows}
    assert list(by_name) == ["Nidia Prueba Uno", "Sin Grupo Prueba", "Maria Prueba Dos"]
    nidia = by_name["Nidia Prueba Uno"]
    assert (nidia.area, nidia.salary_group, nidia.seniority_years, nidia.days_worked, nidia.contribution) == (
        "Departamento de Control",
        12,
        21,
        72,
        3,
    )
    sin = by_name["Sin Grupo Prueba"]
    assert (sin.area, sin.salary_group, sin.seniority_years, sin.contribution) == (
        "Desarrollo Habitacional",
        None,
        None,
        None,
    )
    maria = by_name["Maria Prueba Dos"]
    assert (maria.area, maria.days_worked, maria.salary_group) == (
        "Desarrollo Habitacional",
        60,
        20,
    )  # unnumbered header = sub-group
    assert (nidia.area_number, sin.area_number) == (3, 4)
    assert len(result.errors) == 1 and "#REF" in result.errors[0].message


def test_read_workers_auto_detects_legacy_sheet(tmp_path):
    from fondo_salario.importers.workers import read_workers_auto

    assert len(read_workers_auto(_legacy_workers_workbook(tmp_path)).rows) == 3
    csv = tmp_path / "w.csv"
    csv.write_text("nombre,area\nA,X\n", encoding="utf-8")
    assert len(read_workers_auto(csv).rows) == 1
