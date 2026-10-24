from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from fondo_salario.domain.models import Position, Settings, Worker
from fondo_salario.importers.table import (
    STRUCTURE_COLUMNS,
    STRUCTURE_SHEET,
    WORKERS_COLUMNS,
    WORKERS_SHEET,
    ImportFormatError,
    plan_structure,
    plan_workers,
    read_structure,
    read_workers,
    write_template,
)
from fondo_salario.persistence.sqlite_repo import ProjectStore
from fondo_salario.ui.controller import Controller


def book(tmp_path, workers=(), structure=(), name="datos.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = WORKERS_SHEET
    ws.append(WORKERS_COLUMNS)
    for row in workers:
        ws.append(row)
    ws = wb.create_sheet(STRUCTURE_SHEET)
    ws.append(STRUCTURE_COLUMNS)
    for row in structure:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


ANA = ["Ana Uno", 112233445, "Dir", "Técnico", "XII", 5, 72, 3, "Sí"]


def test_reads_a_valid_workers_table(tmp_path):
    (row,) = read_workers(book(tmp_path, workers=[ANA])).rows
    assert (row.name, row.national_id, row.area, row.position) == ("Ana Uno", "112233445", "Dir", "Técnico")
    assert (row.salary_group, row.seniority_years, row.days_worked, row.contribution, row.eligible) == (
        12,
        5,
        72,
        3,
        True,
    )


def test_missing_columns_are_refused_naming_them(tmp_path):
    path = tmp_path / "w.csv"
    path.write_text("Nombre y apellidos;Área\nAna;Dir\n", encoding="utf-8")
    with pytest.raises(ImportFormatError, match="Faltan estas columnas.*Grupo"):
        read_workers(path)


def test_titles_must_be_in_the_first_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = WORKERS_SHEET
    ws.append(["Listado de personal"])
    ws.append(WORKERS_COLUMNS)
    path = tmp_path / "x.xlsx"
    wb.save(path)
    with pytest.raises(ImportFormatError):
        read_workers(path)


def test_wrong_sheet_name_is_refused_unless_it_is_the_only_one(tmp_path):
    wb = Workbook()
    wb.active.title = "Hoja1"
    wb.create_sheet("Otra")
    path = tmp_path / "x.xlsx"
    wb.save(path)
    with pytest.raises(ImportFormatError, match="no tiene la hoja"):
        read_workers(path)


def test_row_problems_are_reported_with_the_line_and_do_not_stop_the_rest(tmp_path):
    rows = [
        ANA,
        ["", "", "Dir", "", "III", 1, 72, 4, "No"],  # no name
        ["Beto", "", "Dir", "", "ZZZ", 1, 72, 4, "No"],  # bad group
        ["Carla", "", "Dir", "", "III", "x", 72, 4, "No"],  # bad seniority
        ["Dani", "", "Dir", "", "III", 1, 72, 9, "No"],  # contribution out of range
        ["Eva", "", "Dir", "", "III", 1, 72, 4, "quizás"],  # bad yes/no
    ]
    result = read_workers(book(tmp_path, workers=rows))
    assert [r.name for r in result.rows] == ["Ana Uno"]
    assert [e.split(":")[0] for e in result.errors] == ["Fila 3", "Fila 4", "Fila 5", "Fila 6", "Fila 7"]


def test_empty_right_to_collect_means_no_and_blank_rows_are_skipped(tmp_path):
    rows = [["Ana", "", "Dir", "", "XII", 5, 72, 3, None], [None] * 9]
    result = read_workers(book(tmp_path, workers=rows))
    assert [r.eligible for r in result.rows] == [False] and result.errors == []


def test_plan_workers_adds_only_new_ones_and_skips_unknown_areas(tmp_path):
    from fondo_salario.domain.models import Area

    areas = [Area(1, "Dir", 1)]
    existing = [Worker(1, "ana  UNO", "", 1, 20, 9, 1, 1)]
    rows = [
        ANA,
        ["Beto", "999", "dir", "", "III", 1, 72, 4, "No"],
        ["Carla", "", "Zzz", "", "III", 1, 72, 4, "No"],
    ]
    plan = plan_workers(read_workers(book(tmp_path, workers=rows)), areas, existing)
    assert [w.name for w in plan.new] == ["Beto"] and plan.unchanged == 1
    assert "«Zzz»" in plan.skipped[0]


def test_worker_with_the_same_id_is_not_added_again(tmp_path):
    from fondo_salario.domain.models import Area

    areas = [Area(1, "Dir", 1)]
    existing = [Worker(1, "Otro Nombre", "112233445", 1, 12, 5, 72, 3)]
    plan = plan_workers(read_workers(book(tmp_path, workers=[ANA])), areas, existing)
    assert plan.new == [] and plan.unchanged == 1


def test_structure_table_gives_areas_positions_and_empty_areas(tmp_path):
    rows = [
        ["Dir", "Jefe", 2, 1, 1, 1, 6060, 3],
        ["Dir", "Chofer", 1, 1, 0, None, "2 200,50", 2],
        ["Archivo"],
    ]
    result = read_structure(book(tmp_path, structure=rows))
    assert result.errors == []
    jefe, chofer, archivo = result.rows
    assert (jefe.approved, jefe.vacant, jefe.temp_vacant, jefe.distribute) == (2, 1, 1, 1)
    assert chofer.distribute is None
    assert archivo.position == ""


def test_distribute_cannot_exceed_the_vacancies(tmp_path):
    result = read_structure(book(tmp_path, structure=[["Dir", "Jefe", 2, 1, 0, 3, 100, 1]]))
    assert result.rows == [] and "superan las vacantes" in result.errors[0]


def test_plan_structure_adds_what_is_new_and_leaves_existing_alone(tmp_path):
    from fondo_salario.domain.models import Area

    areas = [Area(1, "Dir", 1)]
    positions = [Position(1, 1, "Jefe", 1, 1, 0, D("5000"), 3)]
    rows = [
        ["dir", "jefe", 1, 1, 0, None, 9999, 3],
        ["Dir", "Chofer", 1, 1, 0, None, 100, 1],
        ["Eco", "Contador", 1, 1, 0, None, 1, 1],
    ]
    plan = plan_structure(read_structure(book(tmp_path, structure=rows)), areas, positions)
    assert plan.new_areas == ["Eco"]
    assert [(p.area, p.position) for p in plan.new_positions] == [("Dir", "Chofer"), ("Eco", "Contador")]
    assert plan.unchanged == 1


def test_template_has_only_titles_and_instructions(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(path)
    wb = load_workbook(path)
    assert wb.sheetnames == [WORKERS_SHEET, STRUCTURE_SHEET, "Instrucciones"]
    assert [c.value for c in wb[WORKERS_SHEET][1]] == WORKERS_COLUMNS
    assert wb[WORKERS_SHEET].max_row == 1 and wb[STRUCTURE_SHEET].max_row == 1
    assert read_workers(path).rows == [] and read_structure(path).rows == []


def _fill(ctl):
    ctl.save_settings(Settings(total_amount=D("1000.00"), months_paid=3))
    dir_ = ctl.add_area("Dirección")
    ctl.add_area("Archivo")  # area without positions
    ctl.add_position(Position(0, dir_.id, "Jefe", 2, 1, 1, D("6060.50"), 3, 1))
    ctl.add_position(Position(0, dir_.id, "Chofer", 1, 1, 0, D("2200"), 3))
    ctl.add_worker(Worker(0, "Ana Uno", "00112233445", dir_.id, 12, 5, 72, 3, "Técnico", True))
    ctl.add_worker(Worker(0, "Beto Dos", "", dir_.id, 3, 1, 36, 4, "", False))
    return ctl


def test_exported_report_can_be_imported_into_an_empty_project(tmp_path):
    with ProjectStore(tmp_path / "a.fsal") as a, ProjectStore(tmp_path / "b.fsal") as b:
        source, target = _fill(Controller(a)), Controller(b)
        report = tmp_path / "informe.xlsx"
        source.export(report)

        structure = target.plan_structure_import(report)
        target.apply_structure_import(structure)
        workers = target.plan_workers_import(report)
        assert workers.skipped == [] and structure.skipped == []
        target.apply_workers_import(workers)

        copy, original = target.project, source.project
        assert [a.name for a in copy.areas] == [a.name for a in original.areas]
        assert {
            (p.name, p.approved, p.vacant, p.temp_vacant, p.salary, p.months_unexecuted, p.distribute)
            for p in copy.positions
        } == {
            (p.name, p.approved, p.vacant, p.temp_vacant, p.salary, p.months_unexecuted, p.distribute)
            for p in original.positions
        }
        assert {
            (
                w.full_name,
                w.national_id,
                w.position_name,
                w.salary_group,
                w.seniority_years,
                w.days_worked,
                w.contribution,
                w.eligible,
            )
            for w in copy.workers
        } == {
            (
                w.full_name,
                w.national_id,
                w.position_name,
                w.salary_group,
                w.seniority_years,
                w.days_worked,
                w.contribution,
                w.eligible,
            )
            for w in original.workers
        }

        again = target.plan_workers_import(report)
        assert again.new == [] and again.unchanged == 2  # a second import adds nothing
        assert target.plan_structure_import(report).new_positions == []
