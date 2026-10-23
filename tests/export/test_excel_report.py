from datetime import date
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from fondo_salario.domain.calculation import calculate
from fondo_salario.domain.models import Area, Position, Settings, Worker
from fondo_salario.export.excel_report import export_report, to_roman


def build(tmp_path):
    settings = Settings(total_amount=D("1000.00"), months_paid=3)
    areas = [Area(1, "Direccion", 1), Area(2, "Economia", 2), Area(3, "Sin plazas", 3)]
    positions = [
        Position(1, 1, "Jefe", 1, 1, 0, D("600"), 3),
        Position(2, 2, "Tecnico", 2, 1, 1, D("400"), 3),
    ]
    workers = [
        Worker(1, "Ana Uno Prueba", "00000000001", 1, 12, 10, 72, 4, "Jefe"),
        Worker(2, "Beto Dos Prueba", "00000000002", 2, 20, 3, 72, 3, "Tecnico"),
        Worker(3, "Cira Tres Prueba", "00000000003", 2, 5, 1, 36, 5, "Tecnico"),
        Worker(4, "Dani Cuatro Prueba", "00000000004", 3, 5, 1, 72, 5, "Otro"),
        Worker(5, "Eli Cinco Prueba", "00000000005", 2, 5, 1, 72, 5, "Tecnico", eligible=False),
        Worker(6, "Fabi Seis Prueba", "00000000006", 1, 5, 1, 72, 9, "Jefe"),
    ]
    result = calculate(settings, areas, positions, workers)
    path = tmp_path / "report.xlsx"
    export_report(path, settings, result, generated_on=date(2026, 1, 15))
    return load_workbook(path), result


def test_roman_numerals():
    assert [to_roman(n) for n in (4, 12, 19, 25)] == ["IV", "XII", "XIX", "XXV"]


def test_sheets_exist(tmp_path):
    wb, _ = build(tmp_path)
    assert wb.sheetnames == ["Hoja 1 - Montos por área", "Hoja 2 - Coeficientes", "Pre-nómina", "Parámetros"]


def test_prenomina_lists_every_worker_and_total(tmp_path):
    wb, result = build(tmp_path)
    ws = wb["Pre-nómina"]
    rows = [r for r in ws.iter_rows(values_only=True) if r[1] and str(r[2] or "").isdigit()]
    assert [r[1] for r in rows] == [w.worker.full_name for w in result.workers]
    amounts = {r[1]: r[4] for r in rows}
    assert amounts["Dani Cuatro Prueba"] == 0 and amounts["Eli Cinco Prueba"] == 0
    assert sum(D(str(a)) for a in amounts.values()) == D("1000.00")
    total_row = next(r for r in ws.iter_rows(values_only=True) if r[1] == "TOTAL")
    assert total_row[4] == 1000.0


def test_hoja1_percentages_and_total(tmp_path):
    wb, _ = build(tmp_path)
    ws = wb["Hoja 1 - Montos por área"]
    rows = list(ws.iter_rows(values_only=True))
    grand = next(r for r in rows if r[0] == "TOTAL GENERAL")
    assert grand[9] == 100.0 and grand[10] == 1000.0
    assert any("OK" in str(r[0]) for r in rows if r[0] and "100" in str(r[0]))


def test_hoja2_has_amount_per_worker_and_area_subtotals(tmp_path):
    wb, result = build(tmp_path)
    ws = wb["Hoja 2 - Coeficientes"]
    rows = list(ws.iter_rows(values_only=True))
    ana = next(r for r in rows if r[1] == "Ana Uno Prueba")
    assert ana[-1] == pytest.approx(float(result.workers[0].amount))
    assert any(r[0] and str(r[0]).startswith("Subtotal") for r in rows)


def test_parameters_sheet_lists_weights_and_issues(tmp_path):
    wb, _ = build(tmp_path)
    text = " ".join(str(c) for r in wb["Parámetros"].iter_rows(values_only=True) for c in r if c)
    assert "0.40" in text and "2026-01-15" in text and "datos no válidos" in text.lower()


def test_print_setup(tmp_path):
    wb, _ = build(tmp_path)
    ws = wb["Pre-nómina"]
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1


def test_hoja1_shows_vacancies_to_distribute(tmp_path):
    wb, _ = build(tmp_path)
    rows = list(wb["Hoja 1 - Montos por área"].iter_rows(values_only=True))
    header = next(r for r in rows if r[0] == "Relación de cargos por áreas")
    assert header[4] == "Total de vacantes" and header[5] == "Vacantes a distribuir"
    tecnico = next(r for r in rows if r[0] == "Tecnico")
    assert (tecnico[4], tecnico[5]) == (2, 2)


def test_reports_show_who_has_the_right_to_collect(tmp_path):
    wb, _ = build(tmp_path)
    hoja2 = list(wb["Hoja 2 - Coeficientes"].iter_rows(values_only=True))
    header = next(r for r in hoja2 if r[0] == "Área")
    assert header[4] == "Derecho a cobro"
    assert next(r for r in hoja2 if r[1] == "Ana Uno Prueba")[4] == "Sí"
    assert next(r for r in hoja2 if r[1] == "Eli Cinco Prueba")[4] == "No"
    pre = list(wb["Pre-nómina"].iter_rows(values_only=True))
    assert next(r for r in pre if r[0] == "No.")[5] == "Derecho a cobro"
    assert next(r for r in pre if r[1] == "Eli Cinco Prueba")[5] == "No"
    assert next(r for r in pre if r[1] == "Ana Uno Prueba")[5] == "Sí"
